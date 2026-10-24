#!/usr/bin/env python3
"""金属与世界：保守、可追溯的数据采集脚本。失败时保留历史数据，不编造实时信息。"""
from __future__ import annotations
import hashlib, io, json, os, re, sys, time
import xml.etree.ElementTree as ET
from email.utils import parsedate_to_datetime
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urljoin, urlsplit, urlunsplit

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT/'data'
DATA.mkdir(exist_ok=True)
CN = timezone(timedelta(hours=8))
NOW = datetime.now(timezone.utc)
USER_AGENT = 'MetalWorldResearchDashboard/1.0 (+personal research; public sources only)'
S = requests.Session()
S.headers.update({'User-Agent':USER_AGENT, 'Accept-Language':'en-US,en;q=0.9'})
SOURCES = [
    {'id':'iea-reports','name':'国际能源署（IEA）报告','institution':'IEA','url':'https://www.iea.org/reports', 'host':'iea.org','path':'/reports/', 'tier':'一手机构', 'type':'html'},
    {'id':'iea-news','name':'国际能源署（IEA）新闻','institution':'IEA','url':'https://www.iea.org/news/', 'host':'iea.org','path':'/news/', 'tier':'一手机构', 'type':'html'},
    {'id':'usgs-news','name':'美国地调局（USGS）','institution':'USGS','url':'https://www.usgs.gov/news', 'host':'usgs.gov','path':'/news/', 'tier':'一手机构', 'type':'html'},
    {'id':'icsg','name':'国际铜研究组（ICSG）','institution':'ICSG','url':'https://icsg.org/press-releases/','host':'icsg.org','path':'/download/','tier':'一手机构','type':'html'},
    {'id':'bhp','name':'必和必拓（BHP）','institution':'BHP','url':'https://www.bhp.com/investor-centre/market-announcements','host':'bhp.com','path':'/news/', 'tier':'企业公告','type':'html'},
    {'id':'rio','name':'力拓（Rio Tinto）','institution':'Rio Tinto','url':'https://www.riotinto.com/en/news/releases','host':'riotinto.com','path':'/en/news/releases/', 'tier':'企业公告','type':'html'},
    {'id':'mit-ai','name':'MIT News · 人工智能研究','institution':'MIT News','url':'https://news.mit.edu/topic/mitartificial-intelligence2-rss.xml','host':'news.mit.edu','path':'/','tier':'科研机构','type':'rss','focus':'tech'},
    {'id':'mit-robotics','name':'MIT News · 机器人研究','institution':'MIT News','url':'https://news.mit.edu/topic/mitrobotics-rss.xml','host':'news.mit.edu','path':'/','tier':'科研机构','type':'rss','focus':'tech'},
    {'id':'sia','name':'SIA · 全球半导体产业动态','institution':'SIA','url':'https://www.semiconductors.org/news-events/latest-news/','host':'semiconductors.org','path':'/','tier':'产业协会','type':'html','focus':'tech'},
    {'id':'asml','name':'ASML 半导体设备公告','institution':'ASML','url':'https://investor.asml.com/news/press-releases','host':'investor.asml.com','path':'/news-releases/news-release-details/','tier':'企业公告','type':'html','focus':'tech','always_relevant':True},
    {'id':'tsmc','name':'台积电产业公告（TSMC）','institution':'TSMC','url':'https://pr.tsmc.com/english/latest-news','host':'pr.tsmc.com','path':'/english/news/','tier':'企业公告','type':'html','focus':'tech','always_relevant':True},
    {'id':'mining','name':'MINING.COM行业媒体','institution':'MINING.COM','url':'https://www.mining.com/feed/','host':'mining.com','path':'/','tier':'媒体线索','type':'rss'},
]
OFFICIAL_LINKS = [
 {'name':'世界银行 · Pink Sheet（月度价格）','url':'https://www.worldbank.org/en/research/commodity-markets','frequency':'每月','detail':'提供历史月度名义美元价格，非即时行情'},
 {'name':'IEA · 2026关键矿产展望','url':'https://www.iea.org/reports/global-critical-minerals-outlook-2026','frequency':'年度/专题','detail':'全球关键矿产供需、投资、供应链'},
 {'name':'USGS · 矿产统计','url':'https://www.usgs.gov/centers/national-minerals-information-center/mineral-commodity-summaries','frequency':'年度','detail':'全球主要矿种产量、资源储量'},
 {'name':'ICSG · 铜业统计','url':'https://icsg.org/selected-copper-statistics/','frequency':'月度/预测','detail':'铜矿及精炼铜供需和预测'},
 {'name':'LME · 官方金属价格','url':'https://www.lme.com/en/market-data/reports-and-data/lme-official-prices','frequency':'依许可','detail':'受访问与使用授权限制；本平台不显示未授权实时价格'},
 {'name':'Rio Tinto · 公司公告','url':'https://www.riotinto.com/en/news/releases','frequency':'按公告','detail':'矿山项目和生产经营信息'},
 {'name':'MIT News · AI与机器人','url':'https://news.mit.edu/rss','frequency':'持续发布','detail':'麻省理工学院科研动态，关注AI、机器人、先进材料和能源'},
 {'name':'SIA · 半导体产业资讯','url':'https://www.semiconductors.org/news-events/latest-news/','frequency':'新闻稿/统计','detail':'半导体销售、晶圆制造和国际产业链变化'},
 {'name':'IEA · AI与能源','url':'https://www.iea.org/reports/key-questions-on-energy-and-ai','frequency':'专题/年度','detail':'AI数据中心与能源供应、产业链的交互'},
 {'name':'ASML · 半导体技术及企业公告','url':'https://investor.asml.com/news/press-releases','frequency':'按公告','detail':'半导体制造设备、先进光刻技术及产业合作'},
 {'name':'TSMC · 官方新闻','url':'https://pr.tsmc.com/english/latest-news','frequency':'按公告','detail':'晶圆制造、先进制程及半导体产业投资'},
 {'name':'BHP · 交易所公告','url':'https://www.bhp.com/investor-centre/market-announcements','frequency':'按公告','detail':'业绩、矿山项目及生产经营披露'},
]
METALS = [ ('铜',r'\bcopper\b|\bcu\b|铜'),('铝',r'\balumin(?:ium|um)\b|\bbauxite\b|铝'),('锂',r'\blithium\b|锂'),('镍',r'\bnickel\b|镍'),('钴',r'\bcobalt\b|钴'),('稀土',r'rare earth|rare-earth|\bree\b|稀土'),('锌',r'\bzinc\b|锌'),('锡',r'\btin\b|锡'),('锗镓',r'\bgallium\b|\bgermanium\b|镓|锗'),('关键矿产',r'critical mineral|critical material|strategic mineral|关键矿产') ]
TECH = [
 ('AI与算力',r'artificial intelligence|\bai\b|generative ai|data cent(?:er|re)|machine learning|computing|gpu|人工智能|算力|数据中心'),
 ('半导体',r'semiconductor|chip|foundr|microelectronic|wafer|semiconductor sales|半导体|芯片|晶圆'),
 ('机器人',r'robot|automation|autonomous systems|机器人|自动化'),
 ('新能源与储能',r'batter(?:y|ies)|electric vehicle|\bev\b|renewable|energy storage|solar|hydrogen|电池|储能|新能源'),
 ('先进制造',r'additive manufactur|3d print|advanced manufactur|industrial automation|先进制造|增材制造'),
 ('新材料',r'advanced material|new material|materials science|composite|superconduct|新材料|超导')]
MINING_KEYWORD = re.compile(r'copper|alumin|bauxite|lithium|nickel|cobalt|rare earth|critical mineral|mining|mineral|smelt|refin|zinc|\btin\b|gallium|germanium|tungsten|electric grid|battery|batteries|\bmetal|\bore\b|iron ore|molybdenum|exploration|supply chain',re.I)
TECH_KEYWORD = re.compile(r'artificial intelligence|\bAI\b|machine learning|data cent|semiconductor|chip|foundr|robot|automation|battery|batteries|electric vehicle|energy storage|advanced manufactur|materials science|additive manufactur|quantum comput|solar power|photovoltaic|computing infrastructure|power grid',re.I)
KEYWORD = re.compile(MINING_KEYWORD.pattern+'|'+TECH_KEYWORD.pattern,re.I)


def load(path, default):
    try: return json.loads((DATA/path).read_text(encoding='utf-8'))
    except (FileNotFoundError,ValueError): return default

def save(path,obj):
    (DATA/path).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

def fetch(url,timeout=22):
    resp=S.get(url,timeout=timeout,allow_redirects=True)
    resp.raise_for_status()
    if len(resp.content)>8_000_000: raise ValueError('响应过大（大于8MB）')
    return resp

def clean(x): return re.sub(r'\s+', ' ', x or '').strip()

def normalize(url):
    p=urlsplit(url)
    return urlunsplit((p.scheme,p.netloc.lower(),p.path.rstrip('/') or '/', '', ''))

def parse_date(x):
    if not x: return ''
    text=clean(str(x))
    tests=[r'(20\d{2}-\d{2}-\d{2})',r'(20\d{2}/\d{2}/\d{2})', r'(\d{1,2}[/.-]\d{1,2}[/.-]20\d{2})',r'(\d{1,2} [A-Z][a-z]+ 20\d{2})',r'([A-Z][a-z]+ \d{1,2},? 20\d{2})']
    for regex in tests:
        m=re.search(regex,text)
        if not m: continue
        z=m.group(1)
        try:
            if re.match(r'^20\d{2}-',z): return z
            if re.match(r'^20\d{2}/',z): return z.replace('/','-')
            for fmt in ('%d %B %Y','%b %d, %Y','%B %d, %Y','%B %d %Y','%m/%d/%Y','%d/%m/%Y'):
                try: return datetime.strptime(z,fmt).strftime('%Y-%m-%d')
                except ValueError: pass
        except ValueError: pass
    m=re.search(r'(?<!\d)(\d{1,2})/(\d{1,2})/(\d{2})(?!\d)',text)
    if m:
        try:return datetime.strptime(m.group(0),'%m/%d/%y').strftime('%Y-%m-%d')
        except ValueError:pass
    return ''

def classify(headline, snippet, origin, focus=''):
    txt=(headline+' '+(snippet or ''))[:900]
    metals=[name for name,pattern in METALS if re.search(pattern,txt,re.I)]
    technology=[name for name,pattern in TECH if re.search(pattern,txt,re.I)]
    if focus=='tech' or origin in ('MIT News','SIA','ASML','TSMC'):
        cat='全球科技产业'
    elif origin in ('ICSG','World Bank') or (metals and re.search(r'price|supply|demand|forecast|market|statistic|commodit',txt,re.I)):
        cat='市场数据'
    elif technology: cat='金属与科技'
    else: cat='海外矿业'
    return cat, metals[:4], technology[:3]

def item(title, url, source, date='', snippet=''):
    title=clean(title)
    url=normalize(url)
    if not title or len(title)<13 or len(title)>240 or not url.startswith('https://'): return None
    if not source.get('always_relevant') and not (TECH_KEYWORD if source.get('focus')=='tech' else KEYWORD).search(title+' '+snippet): return None
    cat,metals,tech=classify(title,snippet,source['institution'],source.get('focus',''))
    return {'id':hashlib.sha256(url.encode()).hexdigest()[:16],'title':title,'title_zh':'','url':url,
      'source':source['institution'],'source_name':source['name'],'source_type':source['tier'],
      'published_at':parse_date(date),'fetched_at':NOW.isoformat(timespec='seconds'),'summary':'',
      'original_excerpt':clean(snippet)[:480], 'category':cat,'metals':metals,'tech':tech,
      'verification':'已识别原始来源，中文解读待复核', 'editorial':False}

def article_metadata(url):
    """Obtain date and concise source-authored excerpt from a public article page.

    Date and excerpt are *unverified raw source material*, never an editorial fact-check.
    Request only a few detail pages per source per run.
    """
    resp=fetch(url,timeout=14)
    soup=BeautifulSoup(resp.text,'html.parser')
    meta={}
    for m in soup.find_all('meta'):
        key=(m.get('property') or m.get('name') or '').lower()
        if key and m.get('content'):meta[key]=clean(m['content'])
    ds=[meta.get(k,'') for k in ('article:published_time','datepublished','date','pubdate','article:modified_time')]
    for script in soup.find_all('script',type='application/ld+json')[:5]:
        try:
            obj=json.loads(script.string or script.get_text() or '{}')
            items=obj if isinstance(obj,list) else [obj]
            for i in items:
                if isinstance(i,dict):ds.extend([i.get('datePublished',''),i.get('dateCreated','')])
        except (ValueError,TypeError):pass
    first=soup.find('time')
    if first:ds.extend([first.get('datetime',''),first.get_text(' ',strip=True)])
    pub=next((d for v in ds if isinstance(v,str) and (d:=parse_date(v))), '')
    excerpt=meta.get('description') or meta.get('og:description') or meta.get('twitter:description') or ''
    return pub,clean(excerpt)[:480]


def list_html(src):
    resp=fetch(src['url'])
    soup=BeautifulSoup(resp.text,'html.parser')
    for t in soup(['script','style','footer','nav','noscript']): t.decompose()
    ret=[]
    seen=set()
    for a in soup.find_all('a',href=True):
        url=urljoin(resp.url,a['href'])
        u=urlsplit(url)
        if u.netloc.lower().removeprefix('www.') != src['host'].removeprefix('www.'): continue
        if src['path'] not in u.path or normalize(url)==normalize(src['url']):continue
        if src.get('focus')!='tech' and not re.search(r'20\d{2}|/news/|/reports/|/download/',u.path):continue
        if src.get('focus')=='tech' and len(u.path.strip('/'))<14:continue
        if normalize(url) in seen:continue
        seen.add(normalize(url))
        # anchor text first; avoid parent block repeating numerous unrelated titles
        title=clean(a.get_text(' ',strip=True))
        if len(title)<14 or len(title)>190:
            h=a.find(['h2','h3','h4'])
            title=clean(h.get_text(' ',strip=True)) if h else title
        if len(title)>180:continue
        parent=a.find_parent(['article','li','section']) or a.parent
        ctx=clean(parent.get_text(' ',strip=True))[:500] if parent else ''
        if not src.get('always_relevant') and not (TECH_KEYWORD if src.get('focus')=='tech' else KEYWORD).search(title+' '+ctx):continue
        d=''
        if parent:
            time_el=parent.find('time')
            if time_el: d=time_el.get('datetime') or time_el.get_text(' ',strip=True)
        d=parse_date(d or ctx)
        x=item(title,url,src,d,ctx)
        if x: ret.append(x)
        if len(ret)>=35:break
    # Most publisher index pages do not include article dates in anchor parents.
    # Resolve a few of the most recent links so they can qualify for daily selection.
    detail_limit=6
    visited=0
    for article in ret[:15]:
        if article['published_at'] and len(article.get('original_excerpt',''))>60:continue
        if visited>=detail_limit:break
        visited+=1
        try:
            date,description=article_metadata(article['url'])
            if date and not article['published_at']:article['published_at']=date
            if description and (not article.get('original_excerpt') or len(article['original_excerpt'])<60):
                article['original_excerpt']=description
        except requests.RequestException:pass
        except (ValueError,TypeError):pass
    return ret

def list_rss(src):
    resp=fetch(src['url'])
    root=ET.fromstring(resp.content)
    nodes=list(root.findall('.//item'))
    if not nodes: nodes=list(root.findall('.//{http://www.w3.org/2005/Atom}entry'))
    if not nodes: raise ValueError('RSS/Atom无可用条目')
    def ch(n,names):
        for name in names:
            child=n.find(name)
            if child is not None:
                if child.text:return child.text
                if child.get('href'):return child.get('href')
        return ''
    ret=[]
    for e in nodes[:100]:
        title=ch(e,['title','{http://www.w3.org/2005/Atom}title'])
        link=ch(e,['link','{http://www.w3.org/2005/Atom}link'])
        summary=ch(e,['description','{http://www.w3.org/2005/Atom}summary','{http://purl.org/rss/1.0/modules/content/}encoded'])
        summary=BeautifulSoup(summary or '','html.parser').get_text(' ',strip=True)
        dt=ch(e,['pubDate','{http://www.w3.org/2005/Atom}published','{http://www.w3.org/2005/Atom}updated'])
        try:date=parsedate_to_datetime(dt).strftime('%Y-%m-%d')
        except Exception:date=parse_date(dt)
        x=item(title,link,src,date,summary)
        if x:ret.append(x)
    return ret

def price_data():
    """读取世行官方月度历史工作簿；表头/单位在文件中探测，禁止硬编码价格。"""
    url='https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx'
    wbresp=S.get(url,timeout=40)
    wbresp.raise_for_status()
    if len(wbresp.content)>15_000_000: raise ValueError('工作簿超过15MB')
    book=load_workbook(io.BytesIO(wbresp.content),read_only=True,data_only=True)
    sheet=next((book[s] for s in book.sheetnames if 'monthly price' in s.lower()),book.active)
    rows=list(sheet.iter_rows(values_only=True))
    ix=None
    for i,row in enumerate(rows[:20]):
        vals=[clean(str(v)).lower() for v in row if v is not None]
        if 'copper' in vals and ('nickel' in vals or 'aluminum' in vals):ix=i;break
    if ix is None: raise ValueError('未识别世行工作簿表头，禁止猜测列号')
    header=[clean(str(v)).lower() if v is not None else '' for v in rows[ix]]
    units=[clean(str(v)) if v is not None else '' for v in rows[ix+1]] if ix+1<len(rows) else []
    target={'铜':'copper','铝':'aluminum','镍':'nickel','锌':'zinc','锡':'tin','铅':'lead'}
    found={cn:header.index(en) for cn,en in target.items() if en in header}
    if len(found)<3:raise ValueError('可识别金属价格种类不足，不输出异常结果')
    values=[]
    for row in rows[ix+1:]:
        d=clean(str(row[0])) if row and row[0] is not None else ''
        mt=re.match(r'^(20\d{2})M(\d{2})$',d,re.I)
        if not mt:continue
        month=f'{mt.group(1)}-{mt.group(2)}'
        for metal,ci in found.items():
            if ci>=len(row):continue
            v=row[ci]
            if not isinstance(v,(int,float)) or isinstance(v,bool):continue
            values.append({'metal':metal,'month':month,'value':round(float(v),3),
               'unit':units[ci] if ci<len(units) else '详见原始工作簿',
               'source':'World Bank Pink Sheet','source_url':url})
    months=sorted({v['month'] for v in values})
    result=[v for v in values if v['month'] in set(months[-18:])]
    if not result:raise ValueError('未读取到价格数值')
    return {'frequency':'月度','source_url':url,'source_name':'世界银行 Pink Sheet','retrieved_at':NOW.isoformat(timespec='seconds'),
       'last_period':months[-1],'not_realtime':True,'series':result}

def run():
    old=load('records.json',[])
    index={o['id']:o for o in old if isinstance(o,dict) and 'id' in o}
    runs=[]
    newcount=0
    for src in SOURCES:
        try:
            results=list_rss(src) if src['type']=='rss' else list_html(src)
            if not results:raise ValueError('页面未识别出符合主题的条目（可能已改版）')
            for x in results:
                if x['id'] in index:
                    o=index[x['id']]
                    # Preserve manually reviewed content and its status.
                    for field in ('title_zh','summary','verification','editorial'):
                        if o.get(field):x[field]=o[field]
                    if not x['published_at']:x['published_at']=o.get('published_at','')
                else:newcount+=1
                index[x['id']]=x
            runs.append({'name':src['name'],'source':src['institution'],'state':'success','count':len(results),'message':'成功获取公开标题和来源链接'})
            print(src['id'],'OK',len(results))
        except Exception as e:
            runs.append({'name':src['name'],'source':src['institution'],'state':'failed','count':0,'message':str(e)[:160]})
            print(src['id'],'ERROR',e,file=sys.stderr)
        time.sleep(0.5)
    def sort_key(x):return (x.get('published_at','') or '0000-00-00',x.get('fetched_at',''))
    records=sorted(index.values(),key=sort_key,reverse=True)[:5000]
    save('records.json',records)
    prices=load('prices.json',{'series':[],'not_realtime':True})
    try:
        fresh=price_data()
        prices=fresh
        runs.append({'name':'世界银行月度价格','source':'World Bank','state':'success','count':len(prices['series']),'message':f"最新统计月份 {prices['last_period']}"})
        print('prices OK',prices['last_period'])
    except Exception as e:
        runs.append({'name':'世界银行月度价格','source':'World Bank','state':'failed','count':0,'message':str(e)[:160]})
        print('prices ERROR',e,file=sys.stderr)
    save('prices.json',prices)
    status={'attempted_at':NOW.isoformat(timespec='seconds'),'timezone':'Asia/Singapore','new_records':newcount,
            'record_count':len(records),'sources':runs,'official_links':OFFICIAL_LINKS,
            'note':'本平台按源站发布时间展示；自动新增内容保留英文标题，中文解读需人工复核。价格为世界银行月度价格，非交易报价。'}
    save('status.json',status)
    # Build briefs first, then bundle all four datasets. The previous version
    # rebuilt the dashboard without daily_briefs, leaving the studio stale.
    from daily_briefs import generate_daily_briefs
    briefs=generate_daily_briefs(records,NOW,DATA)
    payload={'records':records,'prices':prices,'status':status,'daily_briefs':briefs}
    (DATA/'dashboard-data.js').write_text('window.METAL_WORLD_DATA = '+json.dumps(payload,ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
    print(f'complete records={len(records)} new={newcount} source_success={sum(s["state"]=="success" for s in runs)}')

if __name__=='__main__':run()
